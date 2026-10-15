"""PowerPoint (python-pptx): a clean 16:9 deck from a SlideDeck (shaped by the Publisher agent)."""
from __future__ import annotations

import io

from .base import BRAND, INK, MUTED, RenderInput


def render_pptx(inp: RenderInput) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    deck = inp.shaped
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    brand, ink, muted = RGBColor(*BRAND), RGBColor(*INK), RGBColor(*MUTED)

    def text(slide, x, y, w, h, value, size, bold=False, color=ink):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = value
        p.font.size, p.font.bold, p.font.color.rgb = Pt(size), bold, color
        return tf

    def bar(slide):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12))
        shape.fill.solid()
        shape.fill.fore_color.rgb = brand
        shape.line.fill.background()

    # title slide
    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(15, 23, 42)
    bg.line.fill.background()
    text(s, 0.9, 2.4, 11.5, 1.8, deck.title, 40, True, RGBColor(255, 255, 255))
    text(s, 0.9, 4.2, 11.5, 0.8, deck.subtitle, 20, False, RGBColor(147, 197, 253))
    text(s, 0.9, 6.6, 11.5, 0.4, "Prepared with BriefDesk · sources on the last slide", 12, False,
         RGBColor(148, 163, 184))

    for i, sl in enumerate(deck.slides, 2):
        s = prs.slides.add_slide(blank)
        bar(s)
        text(s, 0.7, 0.45, 12, 1.0, sl.title, 30, True, ink)
        tf = s.shapes.add_textbox(Inches(0.8), Inches(1.7), Inches(11.8), Inches(5.0)).text_frame
        tf.word_wrap = True
        for j, b in enumerate(sl.bullets[:6]):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.text = "•  " + b
            p.font.size, p.font.color.rgb = Pt(22 if len(sl.bullets) <= 4 else 19), ink
            p.space_after = Pt(14)
        text(s, 12.2, 6.95, 1.0, 0.4, str(i), 11, False, muted)
        if sl.notes:
            s.notes_slide.notes_text_frame.text = sl.notes

    if inp.sources:
        s = prs.slides.add_slide(blank)
        bar(s)
        text(s, 0.7, 0.45, 12, 1.0, "Sources", 30, True, ink)
        tf = s.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.8), Inches(5.4)).text_frame
        tf.word_wrap = True
        for j, src in enumerate(inp.sources[:14]):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.text = f"[{src['n']}] {src['title'][:90]}  —  {src['url'][:80]}"
            p.font.size, p.font.color.rgb = Pt(12), muted
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
